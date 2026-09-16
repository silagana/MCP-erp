"""Importa la hoja "ALUMNOS REGULARES" del Excel que usa hoy la secretaría
(ej. ALUMNOS_2026.xlsx) a la base de MCP-erp. Migración única — de acá en
adelante los pagos/altas/bajas se registran vía WhatsApp/Telegram o directo
en la base, el Excel deja de ser la fuente de verdad (decisión 2026-09-16).

No borra nada: actualiza por coincidencia (UPSERT). Un alumno del Excel se
matchea contra uno ya existente por DNI si lo tiene, si no por
apellido+nombre normalizado — si matchea, actualiza sus cuotas de este año;
si no matchea con nadie, lo da de alta nuevo. Así no se pierde el
histórico que vino de la migración de st-clares-app (cuotas de años
anteriores, movimientos bancarios ya conciliados, etc.).

Nivel de detalle de pagos: SIMPLE — cada mes con monto cargado se crea como
Cuota con estado=pagada directo, sin generar Pago/Imputacion ni tocar
conciliación bancaria (decisión 2026-09-16; las hojas PAGOS <MES> del
Excel, que son el resumen bancario crudo, no se importan en esta pasada).

Alcance: solo la hoja "ALUMNOS REGULARES". La hoja "TALLERES NUEVO SOL" es
un programa aparte con estructura distinta (necesita un campo "grado
escolar" que no existe en el esquema) — queda para una etapa futura.

Requiere la migración previa que permite DNI/CUIT nulos (ver
migrations/versions/ — "allow null dni and cuit").

Uso:
    DATABASE_URL=... python scripts/importar_excel_alumnos.py ALUMNOS_2026.xlsx --dry-run
    DATABASE_URL=... python scripts/importar_excel_alumnos.py ALUMNOS_2026.xlsx
"""
import argparse
import os
import re
import sys
import unicodedata
from datetime import date, datetime
from decimal import Decimal

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import openpyxl
from sqlalchemy import create_engine, func
from sqlalchemy.orm import Session

from app.models import (
    Alumno, Curso, EstadoAlumnoEnum, EstadoCuotaEnum, Inscripcion,
    ModalidadEnum, ReferentePago, Sede, TipoCuotaEnum, VinculoEnum, Cuota,
)

# columna 1-based -> (periodo YYYY-MM, col_monto, col_fecha_o_None)
MESES = [
    ("2026-03", 14, None),
    ("2026-04", 15, 16),
    ("2026-05", 17, 18),
    ("2026-06", 19, 20),
    ("2026-07", 21, 22),
    ("2026-08", 23, 24),
    ("2026-09", 25, 26),
    ("2026-10", 27, 28),
    ("2026-11", 30, 31),
    ("2026-12", 33, 34),
]
COL_ALUMNOS = 1
COL_DNI = 2
COL_SEDE = 3
COL_CURSO = 4
COLS_REFERENTES = [(5, 6), (7, 8), (9, 10), (11, 12)]  # (cuit, nombre)
COL_MATRICULA = 13
COL_DER_EX = 35

SEDES_CONOCIDAS = {"caballito": "Caballito", "boedo": "Boedo", "on line": "Online", "online": "Online"}
MOTIVO_BAJA_IMPORT = "Baja detectada en import de Excel (columna marcada con x)"


def _limpiar_texto(v) -> str | None:
    if v is None:
        return None
    s = str(v)
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("⁠", "").replace("\xa0", " ")
    s = " ".join(s.split())
    return s or None


def _es_x(v) -> bool:
    return isinstance(v, str) and v.strip().lower() == "x"


def _detectar_periodo_baja(valores_por_mes: list[tuple[str, object]]) -> str | None:
    """`valores_por_mes` es [(periodo, valor_celda_monto), ...] en orden
    cronológico. Devuelve el primer período de la RACHA DE 'x' QUE LLEGA
    HASTA EL FINAL de los datos cargados, o None si no hay baja.

    No alcanza con "aparece una x en algún lado" — muchos alumnos tienen
    'x' en los primeros meses porque todavía no estaban inscriptos (se dieron
    de alta más tarde en el año), no porque se dieron de baja. Solo cuenta
    como baja si la 'x' se sostiene hasta el último mes con datos (real o
    'x') — un pago real después de una 'x' significa que no fue baja, solo
    un mes sin cuota generada."""
    ultimo_idx_significativo = None
    for i, (_, v) in enumerate(valores_por_mes):
        if _es_x(v) or (isinstance(v, (int, float)) and v > 0):
            ultimo_idx_significativo = i

    if ultimo_idx_significativo is None:
        return None  # no hay ningún dato cargado todavía

    if not _es_x(valores_por_mes[ultimo_idx_significativo][1]):
        return None  # el último dato es un pago real -> sigue activo

    inicio_baja = valores_por_mes[ultimo_idx_significativo][0]
    for i in range(ultimo_idx_significativo, -1, -1):
        periodo_i, valor_i = valores_por_mes[i]
        if _es_x(valor_i):
            inicio_baja = periodo_i
        else:
            break
    return inicio_baja


# Nombres de pila comunes que, cuando aparecen ANTES de la última palabra,
# casi siempre son la primera mitad de un nombre compuesto (ej. "Gentile
# Maria Giuliana" -> apellido "Gentile", nombre "Maria Giuliana") y no parte
# de un apellido compuesto. Encontrado revisando a mano los 54 casos de
# 3+ palabras del padrón real — no es exhaustivo, sigue siendo una
# heurística, no una regla.
_NOMBRES_COMPUESTOS_COMUNES = {
    "maria", "jose", "juan", "ana", "luis", "carlos", "martin", "lautaro",
}


def _parsear_nombre(nombre_completo: str) -> tuple[str, str]:
    """'Apellido[ Apellido2] Nombre[ Nombre2]' -> (apellido, nombre). Asume
    que la ÚLTIMA palabra es el nombre de pila, salvo que la anteúltima sea
    un nombre de pila común (ver _NOMBRES_COMPUESTOS_COMUNES) — ahí toma
    las últimas DOS palabras como nombre compuesto. Heurística razonable
    para apellidos compuestos, pero no perfecta. Los casos de 3+ palabras
    se listan aparte al final del import para revisión manual igual."""
    partes = nombre_completo.split()
    if len(partes) == 1:
        return partes[0], ""
    if len(partes) >= 3:
        anteultima = unicodedata.normalize("NFKD", partes[-2]).encode("ascii", "ignore").decode("ascii").lower()
        if anteultima in _NOMBRES_COMPUESTOS_COMUNES:
            return " ".join(partes[:-2]), " ".join(partes[-2:])
    return " ".join(partes[:-1]), partes[-1]


def _normalizar_clave(apellido: str, nombre: str) -> str:
    s = f"{apellido} {nombre}".lower()
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode("ascii")
    return " ".join(s.split())


def _normalizar_sede(nombre_sede: str) -> str:
    clave = " ".join(nombre_sede.strip().lower().split())
    return SEDES_CONOCIDAS.get(clave, nombre_sede.strip().title())


def _to_decimal(v) -> Decimal | None:
    if v is None or _es_x(v):
        return None
    if isinstance(v, (int, float)) and v > 0:
        return Decimal(str(v))
    return None


def _to_fecha(v) -> date | None:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    return None


def _obtener_o_crear_sede(session: Session, nombre: str) -> Sede:
    nombre_norm = _normalizar_sede(nombre)
    sede = session.query(Sede).filter(func.lower(Sede.nombre) == nombre_norm.lower()).first()
    if sede is None:
        sede = Sede(nombre=nombre_norm, dia_vencimiento_default=10, activa=True)
        session.add(sede)
        session.flush()
    return sede


def _obtener_o_crear_curso(session: Session, nombre: str, sede: Sede, monto_referencia: Decimal | None) -> Curso:
    nombre_norm = nombre.strip().title()
    curso = (
        session.query(Curso)
        .filter(Curso.sede_id == sede.id, func.lower(Curso.nombre) == nombre_norm.lower())
        .first()
    )
    if curso is None:
        base = monto_referencia or Decimal("100000")
        curso = Curso(
            sede_id=sede.id,
            nombre=nombre_norm,
            modalidad=ModalidadEnum.presencial,
            monto_matricula=base,
            monto_cuota_mensual=base,
            cantidad_cuotas=10,
            activo=True,
        )
        session.add(curso)
        session.flush()
    return curso


def _buscar_alumno_existente(session: Session, dni: str | None, apellido: str, nombre: str) -> Alumno | None:
    if dni:
        alumno = session.query(Alumno).filter(Alumno.dni == dni).first()
        if alumno:
            return alumno
    clave = _normalizar_clave(apellido, nombre)
    for candidato in session.query(Alumno).all():
        if _normalizar_clave(candidato.apellido, candidato.nombre) == clave:
            return candidato
    return None


def importar(xlsx_path: str, database_url: str, dry_run: bool) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb["ALUMNOS REGULARES"]

    engine = create_engine(database_url, future=True)
    with Session(engine) as session:
        nuevos = actualizados = omitidos = 0
        nombres_revisar: list[str] = []
        cuotas_creadas = cuotas_actualizadas = 0

        for fila in range(2, ws.max_row + 1):
            nombre_crudo = _limpiar_texto(ws.cell(row=fila, column=COL_ALUMNOS).value)
            curso_crudo = _limpiar_texto(ws.cell(row=fila, column=COL_CURSO).value)
            sede_cruda = _limpiar_texto(ws.cell(row=fila, column=COL_SEDE).value)

            if not nombre_crudo or not curso_crudo or _es_x(curso_crudo) or not sede_cruda:
                omitidos += 1
                continue

            apellido, nombre = _parsear_nombre(nombre_crudo)
            if len(nombre_crudo.split()) >= 3:
                nombres_revisar.append(nombre_crudo)

            dni_val = ws.cell(row=fila, column=COL_DNI).value
            dni = str(int(dni_val)) if isinstance(dni_val, (int, float)) else (_limpiar_texto(dni_val))

            sede = _obtener_o_crear_sede(session, sede_cruda)

            alumno = _buscar_alumno_existente(session, dni, apellido, nombre)
            es_nuevo = alumno is None
            if es_nuevo:
                alumno = Alumno(nombre=nombre, apellido=apellido, dni=dni, sede_id=sede.id, estado=EstadoAlumnoEnum.activo)
                session.add(alumno)

            alumno.sede_id = sede.id
            if dni and not alumno.dni:
                alumno.dni = dni
            # flush ya acá: de acá en más siempre usamos alumno.id (columna),
            # nunca la relación `alumno=` — mezclar las dos en un mismo insert
            # hace que SQLAlchemy pise el alumno_id con None al hacer flush.
            session.flush()

            # referentes de pago: agrega los que falten (matcheados por nombre)
            for col_cuit, col_nombre in COLS_REFERENTES:
                nombre_ref = _limpiar_texto(ws.cell(row=fila, column=col_nombre).value)
                if not nombre_ref:
                    continue
                cuit_val = ws.cell(row=fila, column=col_cuit).value
                cuit = str(int(cuit_val)) if isinstance(cuit_val, (int, float)) else _limpiar_texto(cuit_val)
                ya_existe = any(
                    _limpiar_texto(r.nombre_completo) == nombre_ref for r in (alumno.referentes or [])
                )
                if not ya_existe:
                    session.add(ReferentePago(
                        alumno_id=alumno.id,
                        nombre_completo=nombre_ref,
                        cuit_cuil=cuit,
                        vinculo=VinculoEnum.otro,
                        es_default=(col_cuit == 5),
                    ))

            # detectar mes de baja (racha de 'x' que llega hasta el final)
            valores_meses = [(periodo, ws.cell(row=fila, column=col_monto).value) for periodo, col_monto, _ in MESES]
            periodo_baja = _detectar_periodo_baja(valores_meses)

            monto_ref_curso = None
            for _, col_monto, _ in MESES:
                d = _to_decimal(ws.cell(row=fila, column=col_monto).value)
                if d:
                    monto_ref_curso = d
                    break

            curso = _obtener_o_crear_curso(session, curso_crudo, sede, monto_ref_curso)

            inscripcion = (
                session.query(Inscripcion)
                .filter(Inscripcion.alumno_id == alumno.id, Inscripcion.curso_id == curso.id)
                .first()
            )
            if inscripcion is None:
                inscripcion = Inscripcion(
                    alumno_id=alumno.id,
                    curso_id=curso.id,
                    fecha_inscripcion=date(2026, 3, 1),
                    activa=(periodo_baja is None),
                )
                session.add(inscripcion)
                session.flush()
            else:
                inscripcion.activa = periodo_baja is None

            if periodo_baja and alumno.estado == EstadoAlumnoEnum.activo:
                alumno.estado = EstadoAlumnoEnum.baja
                anio, mes = periodo_baja.split("-")
                alumno.fecha_estado = date(int(anio), int(mes), 1)
                alumno.motivo_estado = MOTIVO_BAJA_IMPORT
            elif (
                periodo_baja is None
                and alumno.estado == EstadoAlumnoEnum.baja
                and alumno.motivo_estado == MOTIVO_BAJA_IMPORT
            ):
                # re-corrida con la detección de baja arreglada (2026-09-16):
                # este alumno había quedado marcado de baja por error (la
                # detección vieja tomaba cualquier 'x' como corte definitivo,
                # aunque después hubiera pagos reales) — se revierte.
                alumno.estado = EstadoAlumnoEnum.activo
                alumno.fecha_estado = None
                alumno.motivo_estado = None

            session.flush()  # asegura alumno.id/inscripcion.id para lo que sigue

            # matrícula
            monto_matricula = _to_decimal(ws.cell(row=fila, column=COL_MATRICULA).value)
            if monto_matricula:
                c, accion = _upsert_cuota(session, inscripcion.id, TipoCuotaEnum.matricula, None, monto_matricula, date(2026, 3, 1))
                if accion == "creada":
                    cuotas_creadas += 1
                else:
                    cuotas_actualizadas += 1

            # cuotas mensuales
            for periodo, col_monto, col_fecha in MESES:
                if periodo_baja and periodo >= periodo_baja:
                    continue
                monto = _to_decimal(ws.cell(row=fila, column=col_monto).value)
                if not monto:
                    continue
                fecha = _to_fecha(ws.cell(row=fila, column=col_fecha).value) if col_fecha else None
                anio, mes = periodo.split("-")
                if fecha is None:
                    fecha = date(int(anio), int(mes), sede.dia_vencimiento_default)
                numero = int(mes) - 2  # marzo=1, abril=2, ...
                c, accion = _upsert_cuota(session, inscripcion.id, TipoCuotaEnum.mensual, periodo, monto, fecha, numero_cuota=numero)
                if accion == "creada":
                    cuotas_creadas += 1
                else:
                    cuotas_actualizadas += 1

            # derecho de examen
            monto_der_ex = _to_decimal(ws.cell(row=fila, column=COL_DER_EX).value)
            if monto_der_ex:
                c, accion = _upsert_cuota(session, inscripcion.id, TipoCuotaEnum.examen, "2026-12", monto_der_ex, date(2026, 12, 1))
                if accion == "creada":
                    cuotas_creadas += 1
                else:
                    cuotas_actualizadas += 1

            if es_nuevo:
                nuevos += 1
            else:
                actualizados += 1

        print(f"Alumnos nuevos: {nuevos}")
        print(f"Alumnos actualizados: {actualizados}")
        print(f"Filas omitidas (sin nombre/curso/sede válido): {omitidos}")
        print(f"Cuotas creadas: {cuotas_creadas}")
        print(f"Cuotas actualizadas: {cuotas_actualizadas}")
        if nombres_revisar:
            print(f"\nNombres con 3+ palabras (revisar apellido/nombre a mano, {len(nombres_revisar)}):")
            for n in nombres_revisar:
                print(f"  - {n}")

        if dry_run:
            session.rollback()
            print("\n[dry-run] No se guardó nada.")
        else:
            session.commit()
            print("\nGuardado.")


def _upsert_cuota(
    session: Session,
    inscripcion_id: int,
    tipo: TipoCuotaEnum,
    periodo: str | None,
    monto: Decimal,
    fecha_vencimiento: date,
    numero_cuota: int | None = None,
) -> tuple[Cuota, str]:
    query = session.query(Cuota).filter(Cuota.inscripcion_id == inscripcion_id, Cuota.tipo == tipo)
    query = query.filter(Cuota.periodo == periodo) if periodo else query.filter(Cuota.periodo.is_(None))
    cuota = query.first()
    if cuota is None:
        cuota = Cuota(
            inscripcion_id=inscripcion_id,
            tipo=tipo,
            numero_cuota=numero_cuota,
            periodo=periodo,
            fecha_vencimiento=fecha_vencimiento,
            monto_original=monto,
            monto_actualizado=monto,
            estado=EstadoCuotaEnum.pagada,
            saldo_pendiente=Decimal("0.00"),
        )
        session.add(cuota)
        return cuota, "creada"
    cuota.monto_actualizado = monto
    cuota.estado = EstadoCuotaEnum.pagada
    cuota.saldo_pendiente = Decimal("0.00")
    return cuota, "actualizada"


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("xlsx_path", help="Ruta al archivo Excel (hoja 'ALUMNOS REGULARES')")
    parser.add_argument("--dry-run", action="store_true", help="Solo mostrar el resumen, no guardar nada")
    args = parser.parse_args()

    database_url = os.environ.get("DATABASE_URL")
    if not database_url:
        print("Falta DATABASE_URL en el entorno.", file=sys.stderr)
        sys.exit(1)

    importar(args.xlsx_path, database_url, dry_run=args.dry_run)
